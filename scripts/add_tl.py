"""Add the techno-legal classification to the catalogue files in the user folder.
usage: add_tl.py FOLDER TL_OUT_DIR"""
import json, sys, glob, os, csv, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

folder, tl_dir = sys.argv[1:3]
rd = os.path.join(folder, "rights")
rights = json.load(open(os.path.join(rd, "rights_catalogue.json")))
tl = {}
for f in sorted(glob.glob(os.path.join(tl_dir, "*.json"))):
    for x in json.load(open(f)):
        tl[x["id"]] = x
assert len(tl) == len(rights) and all(r["id"] in tl for r in rights)

PAPER = {"R0233": "CLICK", "R0242": "READ", "R0240": "READ", "R0243": "READ", "R0244": "READ", "R0249": "OPERATE",
         "R0250": "OPERATE", "R0251": "CLICK", "R0252": "CLICK", "R0253": "DELEGATE", "R0256": "READ", "R0275": "DELEGATE",
         "R0777": "READ", "R0778": "CLICK", "R0781": "READ", "R0791": "CLICK", "R0801": "CLICK", "R0814": "DELEGATE",
         "R0957": "CLICK", "R0967": "CLICK", "R0968": "CLICK", "R0969": "CLICK", "R0970": "OPERATE", "R0973": "OPERATE",
         "R0975": "OPERATE", "R0981": "BUILD", "R0983": "BUILD", "R1382": "READ", "R1387": "OPERATE", "R1414": "OPERATE",
         "R1416": "OPERATE", "R1623": "BUILD", "R1646": "READ", "R1684": "OPERATE", "R0594": "OPERATE", "R0595": "OPERATE",
         "R1064": "BUILD"}
NEW = ["techno_legal", "tl_level", "tl_family", "tl_scope", "tl_reason", "paper_table2_level"]
for r in rights:
    x = tl[r["id"]]
    core = x["tl"] in ("techno-legal", "borderline")
    r.update({"techno_legal": x["tl"], "tl_level": x["level"] if core else "", "tl_family": x["family"] if core else "",
              "tl_scope": x["scope"], "tl_reason": x["reason"], "paper_table2_level": PAPER.get(r["id"], "")})

json.dump(rights, open(os.path.join(rd, "rights_catalogue.json"), "w"), ensure_ascii=False, indent=1)

with open(os.path.join(rd, "rights_catalogue.csv"), encoding="utf-8-sig") as f:
    head = next(csv.reader(f))
head = [h for h in head if h not in NEW] + NEW
with open(os.path.join(rd, "rights_catalogue.csv"), "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f)
    w.writerow(head)
    for r in rights:
        w.writerow([r.get(h, "") for h in head])

HDR = PatternFill("solid", fgColor="1F3B5C")
TLF = PatternFill("solid", fgColor="B73A3A")
by_id = {r["id"]: r for r in rights}
wb = openpyxl.load_workbook(os.path.join(rd, "rights_catalogue.xlsx"))
for name in ["Applicable law", "In negotiation", "Planned initiative", "All rights"]:
    s = wb[name]
    hdr = [c.value for c in s[1]]
    if "techno_legal" in hdr:
        continue
    start = len(hdr) + 1
    for j, k in enumerate(NEW):
        c = s.cell(row=1, column=start + j, value=k)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = TLF; c.alignment = Alignment(wrap_text=True, vertical="top")
        s.column_dimensions[get_column_letter(start + j)].width = 50 if k == "tl_reason" else 16
    for i in range(2, s.max_row + 1):
        r = by_id[s.cell(row=i, column=1).value]
        for j, k in enumerate(NEW):
            c = s.cell(row=i, column=start + j, value=r[k])
            c.alignment = Alignment(wrap_text=True, vertical="top")
    s.auto_filter.ref = s.dimensions

if "Techno-legal" in wb.sheetnames:
    del wb["Techno-legal"]
s = wb.create_sheet("Techno-legal", 1)
cols = [("id", 8), ("table_status_label", 16), ("act", 30), ("article", 14), ("right", 46), ("rightholder", 24),
        ("techno_legal", 13), ("tl_level", 11), ("tl_family", 22), ("tl_scope", 11), ("tl_reason", 60),
        ("paper_table2_level", 12), ("in_effect_today", 16), ("quote", 60), ("eurlex", 30)]
for j, (k, wdt) in enumerate(cols, 1):
    c = s.cell(row=1, column=j, value=k)
    c.font = Font(bold=True, color="FFFFFF"); c.fill = TLF; c.alignment = Alignment(wrap_text=True, vertical="top")
    s.column_dimensions[get_column_letter(j)].width = wdt
LV = {"CLICK": 0, "READ": 1, "OPERATE": 2, "BUILD": 3, "DELEGATE": 4, "": 5}
sel = [r for r in rights if r["techno_legal"] in ("techno-legal", "borderline")]
sel.sort(key=lambda r: (r["techno_legal"] != "techno-legal", LV[r["tl_level"]]))
for i, r in enumerate(sel, 2):
    for j, (k, _) in enumerate(cols, 1):
        s.cell(row=i, column=j, value=r[k]).alignment = Alignment(wrap_text=True, vertical="top")
s.freeze_panes = "E2"; s.auto_filter.ref = s.dimensions

rm = wb["README"]
if not any(c.value == "techno_legal" for c in rm["A"]):
    rows = [
        ("techno_legal", "Classification under the definition in 'Techno-Legal Rights and Who Can Use Them': techno-legal / borderline / digitized / no. "
                         "Applied literally, including the test 'if the mechanism disappeared, could the right still be exercised, and its result still be used, without technical means?'"),
        ("tl_level", "Gradient level from the paper: CLICK, READ, OPERATE, BUILD, DELEGATE (techno-legal and borderline rows only)"),
        ("tl_family", "Family of the right (the paper's seven families plus complaint & redress, device & software choice, delegation, other)"),
        ("tl_scope", "everyday = natural persons in everyday digital life; specialist; business"),
        ("paper_table2_level", "Level the paper's Table 2 assigns, for the 37 catalogue rows that match its 30 rows"),
    ]
    for k, v in rows:
        rm.append([k, v])

wb.save(os.path.join(rd, "rights_catalogue.xlsx"))

C = collections.Counter(r["techno_legal"] for r in rights)
print("classes", dict(C))
print("levels (techno-legal)", dict(collections.Counter(r["tl_level"] for r in rights if r["techno_legal"] == "techno-legal")))
print("everyday techno-legal", sum(1 for r in rights if r["techno_legal"] == "techno-legal" and r["tl_scope"] == "everyday"))
