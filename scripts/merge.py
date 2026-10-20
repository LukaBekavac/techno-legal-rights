"""Merge agent outputs, verify quotes/articles against the source texts, and write the rights catalogue.

usage: python3 -I merge.py inventory.json RIGHTS_DIR EXTRACT_DIR OUT_DIR
"""
import json, os, re, sys, glob, csv, unicodedata
from collections import Counter, defaultdict

inv = json.load(open(sys.argv[1]))
RD, XD, OUT = sys.argv[2], sys.argv[3], sys.argv[4]
os.makedirs(OUT, exist_ok=True)

STATUS_LABEL = {"applicable": "Applicable law", "in_negotiation": "In negotiation", "planned": "Planned initiative"}
STATUS_ORDER = ["applicable", "in_negotiation", "planned"]
WITHDRAWN = {"52017PC0010", "52023PC0232", "52022PC0496"}
INSTR = {"R": "Regulation", "L": "Directive", "D": "Decision", "PC": "Proposal"}

by_celex = {}
for e in inv:
    for c in e["extract_from"]:
        by_celex.setdefault(c, e)


def instrument(c):
    return "Proposal" if c[5:7] == "PC" else INSTR.get(c[5], c[5])


def legal_status(c, e):
    if c in WITHDRAWN:
        return "Withdrawn proposal (never became law)"
    if c[0] == "5":
        return "Pending proposal"
    if c == "32024L2853":
        return "Adopted, applies to products placed on the market from 9 Dec 2026"
    if c == "32025L2647":
        return "Adopted, transposition period running"
    if c in {"31985L0374"}:
        return "In force (products placed on the market before 9 Dec 2026)"
    return "In force"


def norm(s):
    s = unicodedata.normalize("NFKC", s)
    s = s.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"').replace("–", "-").replace("—", "-")
    s = re.sub(r"▼\w+", " ", s)
    return re.sub(r"\s+", " ", s).strip().lower()


texts, arts = {}, {}
for c in by_celex:
    texts[c] = norm(open(os.path.join(XD, f"{c}.articles.txt"), encoding="utf-8").read())
    arts[c] = set(re.findall(r"^Art\s+(\S+)\s*\|", open(os.path.join(XD, f"{c}.index.txt")).read(), re.M))

rights, sources, problems = [], [], []
for f in sorted(glob.glob(os.path.join(RD, "G*.json"))):
    d = json.load(open(f))
    g = d.get("group", os.path.basename(f)[:3])
    for s in d["sources"]:
        s["group"] = g
        sources.append(s)
    for r in d["rights"]:
        r["group"] = g
        rights.append(r)

seen_src = {s["celex"] for s in sources}
missing_src = sorted(set(by_celex) - seen_src)

for r in rights:
    c = r.get("celex", "")
    e = by_celex.get(c)
    if not e:
        problems.append(f"unknown celex {c} in {r.get('group')}: {r.get('right')}")
        continue
    r["entry_no"] = e["no"]
    r["table_status"] = e["status_in_table"]
    r["table_status_label"] = STATUS_LABEL[e["status_in_table"]]
    r["policy_area"] = e["category"]
    r["table_entry"] = e["table_label"]
    r["table_refs"] = e["table_refs"]
    r["entry_current_status"] = e["current_status"]
    r["instrument"] = instrument(c)
    r["source_legal_status"] = legal_status(c, e)
    r["eurlex"] = f"https://eur-lex.europa.eu/legal-content/EN/TXT/?uri=CELEX:{c}"
    q = norm(r.get("quote", ""))
    q = q.strip(" .…'\"")
    parts = [p.strip() for p in re.split(r"\.\.\.|…|\[\.\.\.\]", q) if len(p.strip()) > 8]
    r["quote_verified"] = bool(parts) and all(p in texts[c] for p in parts)
    m = re.match(r"\s*Arts?\.?\s*(\d+[a-z]*)", r.get("article", ""))
    r["article_verified"] = bool(m and m.group(1) in arts[c]) or "inserted" in r.get("article", "").lower() or " of " in r.get("article", "")

cols = ["table_status_label", "policy_area", "entry_no", "table_entry", "act", "celex", "instrument", "source_legal_status",
        "article", "right", "category", "rightholder", "holder_kind", "addressee", "description", "exercise", "explicit",
        "quote", "applies_from", "notes", "quote_verified", "article_verified", "entry_current_status", "eurlex"]

key = lambda r: (STATUS_ORDER.index(r["table_status"]), r["entry_no"], r["celex"],
                 [int(x) if x.isdigit() else x for x in re.findall(r"\d+|[a-z]+", r.get("article", ""))])
rights = [r for r in rights if "table_status" in r]
rights.sort(key=key)
for i, r in enumerate(rights, 1):
    r["id"] = f"R{i:04d}"
cols = ["id"] + cols

json.dump(rights, open(os.path.join(OUT, "rights_catalogue.json"), "w"), indent=1, ensure_ascii=False)
with open(os.path.join(OUT, "rights_catalogue.csv"), "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.DictWriter(fh, fieldnames=cols, extrasaction="ignore")
    w.writeheader()
    w.writerows(rights)

# ---------- per-law summary
src_by_celex = {s["celex"]: s for s in sources}
cnt = Counter(r["celex"] for r in rights)
laws = []
for e in inv:
    if not e["extract_from"]:
        laws.append(dict(no=e["no"], status=STATUS_LABEL[e["status_in_table"]], area=e["category"], entry=e["table_label"], refs=e["table_refs"],
                         current=e["current_status"], celex="", act="", rights=0, comment="No legal text exists yet: nothing to extract"))
    for c in e["extract_from"]:
        s = src_by_celex.get(c, {})
        laws.append(dict(no=e["no"], status=STATUS_LABEL[e["status_in_table"]], area=e["category"], entry=e["table_label"], refs=e["table_refs"],
                         current=e["current_status"], celex=c, act=s.get("act", ""), rights=cnt.get(c, 0), comment=s.get("comment", "NOT PROCESSED")))

# ---------- Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "README"
readme = [
    ["Rights catalogue: EU digital legislation (Bruegel / kaizenner.eu Table 1)"],
    [""],
    ["Source table", "Tables_digital_eu_legislation.pdf, Table 1 'Overview of EU Legislation in the Digital Sector' (mid-2023)"],
    ["Grouping", "Applicable law / In negotiation / Planned initiative, as shown by the cell colour in the table (status in mid-2023)"],
    ["Current status", "Checked 9 Oct 2026 against the EP Legislative Observatory (OEIL) and the Publications Office (Cellar). See the 'Laws' sheet."],
    ["Texts used", "In-force acts: latest consolidated version; files adopted since 2023: the final adopted act; pending/withdrawn: the Commission proposal"],
    ["What counts as a right", "Something a natural or legal person can actively exercise or claim: express rights, actionable options built into duties (e.g. DSA Art 38), individual information owed to them, and remedies."],
    ["holder_kind", "individual = natural persons only; any person = natural or legal persons; organisation = businesses/legal persons only"],
    ["explicit", "TRUE = the text uses 'right/entitled/may request'; FALSE = derived from another party's duty"],
    ["quote_verified", "TRUE = the quote was found verbatim in the source text (automatic check)"],
    ["Caveat", "Extraction was done by AI and is not legal advice. For directives the right reaches people through national transposition law. Proposal texts may change."],
    [""],
    ["Totals"],
]
for s in STATUS_ORDER:
    readme.append([STATUS_LABEL[s], f"{sum(1 for r in rights if r['table_status']==s)} rights"])
readme.append(["All", f"{len(rights)} rights from {len(set(r['celex'] for r in rights))} legal texts"])
for row in readme:
    ws.append(row)
ws["A1"].font = Font(bold=True, size=14)
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 140

HDR = PatternFill("solid", fgColor="1F3B5C")
FILL = {"applicable": "DCE6F2", "in_negotiation": "E3F4FA", "planned": "F2FBFD"}
widths = {"id": 8, "table_status_label": 16, "policy_area": 18, "entry_no": 6, "table_entry": 30, "act": 16, "celex": 13, "instrument": 11,
          "source_legal_status": 22, "article": 12, "right": 38, "category": 24, "rightholder": 22, "holder_kind": 12, "addressee": 22,
          "description": 60, "exercise": 40, "explicit": 8, "quote": 60, "applies_from": 11, "notes": 30, "quote_verified": 9,
          "article_verified": 9, "entry_current_status": 40, "eurlex": 30}


def sheet(title, rows):
    s = wb.create_sheet(title)
    s.append(cols)
    for c in s[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        s.append([r.get(k, "") for k in cols])
    for i, k in enumerate(cols, 1):
        s.column_dimensions[get_column_letter(i)].width = widths.get(k, 15)
    for row in s.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    s.freeze_panes = "C2"
    s.auto_filter.ref = s.dimensions
    return s


for st in STATUS_ORDER:
    sheet(STATUS_LABEL[st], [r for r in rights if r["table_status"] == st])
sheet("All rights", rights)

s = wb.create_sheet("Laws")
lc = ["no", "status", "area", "entry", "refs", "current", "celex", "act", "rights", "comment"]
s.append(["Table #", "Status in table", "Policy area", "Table entry", "Table references", "Current status (9 Oct 2026)", "Text analysed (CELEX)", "Act", "Rights found", "Extraction comment"])
for c in s[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical="top")
for l in laws:
    s.append([l[k] for k in lc])
for i, wdt in enumerate([7, 16, 20, 40, 26, 50, 14, 18, 8, 60], 1):
    s.column_dimensions[get_column_letter(i)].width = wdt
for row in s.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
s.freeze_panes = "B2"; s.auto_filter.ref = s.dimensions

# pivot: category x status
s = wb.create_sheet("By category")
cats = sorted(Counter(r["category"] for r in rights).items(), key=lambda x: -x[1])
s.append(["Category"] + [STATUS_LABEL[x] for x in STATUS_ORDER] + ["Total"])
for cat, n in cats:
    s.append([cat] + [sum(1 for r in rights if r["category"] == cat and r["table_status"] == st) for st in STATUS_ORDER] + [n])
for c in s[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR
s.column_dimensions["A"].width = 42
wb.save(os.path.join(OUT, "rights_catalogue.xlsx"))

# ---------- Markdown
md = ["# Rights under EU digital legislation", "",
      "Extracted from the laws listed in *Table 1: Overview of EU Legislation in the Digital Sector* (Bruegel / kaizenner.eu, mid-2023). "
      "Grouped by the table's status (cell colour): **Applicable law**, **In negotiation**, **Planned initiative**. "
      "Each act's current status was checked on 9 Oct 2026. AI-assisted extraction, not legal advice.", "",
      "| Group | Table entries | Legal texts analysed | Rights |", "|---|---|---|---|"]
for st in STATUS_ORDER:
    md.append(f"| {STATUS_LABEL[st]} | {sum(1 for e in inv if e['status_in_table']==st)} | "
              f"{len({c for e in inv if e['status_in_table']==st for c in e['extract_from']})} | {sum(1 for r in rights if r['table_status']==st)} |")
md.append(f"| **Total** | {len(inv)} | {len(by_celex)} | {len(rights)} |")
md.append("")
for st in STATUS_ORDER:
    md += [f"## {STATUS_LABEL[st]}", ""]
    for e in [e for e in inv if e["status_in_table"] == st]:
        md.append(f"### #{e['no']} {e['table_label']} ({e['table_refs']}), {e['category']}")
        md.append(f"*Current status:* {e['current_status']}")
        md.append("")
        if not e["extract_from"]:
            md += ["_No legal text yet: nothing to extract._", ""]
            continue
        for c in e["extract_from"]:
            rs = [r for r in rights if r["celex"] == c]
            src = src_by_celex.get(c, {})
            md.append(f"**{src.get('act') or c}** · CELEX [{c}](https://eur-lex.europa.eu/legal-content/EN/TXT/?uri=CELEX:{c}) · {legal_status(c, e)} · {len(rs)} rights")
            md.append("")
            if not rs:
                md += [f"_No person-exercisable rights found. {src.get('comment', '')}_", ""]
                continue
            md += ["| Article | Right | Holder | Category | What it lets you do |", "|---|---|---|---|---|"]
            for r in rs:
                desc = (r.get("description", "") or "").replace("|", "/").replace("\n", " ")
                md.append(f"| {r.get('article','')} | {r.get('right','').replace('|','/')} | {r.get('rightholder','').replace('|','/')} | {r.get('category','')} | {desc} |")
            md.append("")
open(os.path.join(OUT, "RIGHTS.md"), "w").write("\n".join(md))

print("rights:", len(rights), "| by status:", Counter(r["table_status"] for r in rights))
print("sources reported:", len(seen_src), "/", len(by_celex), "missing:", missing_src)
print("quote verified:", sum(r["quote_verified"] for r in rights), "/", len(rights), "| article verified:", sum(r["article_verified"] for r in rights))
print("categories:", Counter(r["category"] for r in rights).most_common())
print("holder_kind:", Counter(r.get("holder_kind") for r in rights))
for p in problems[:20]:
    print("PROBLEM", p)
